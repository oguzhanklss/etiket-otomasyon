#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
COP-31 Demirbaş Etiket Yazdırma
Excel'den okur, Zebra ZD220 (203 dpi) için ZPL üretir, yazıcıya RAW gönderir.

Etiket düzeni (50 x 30 mm):
    ŞARTNAME NO:3
    S/N:35CG6370JV6
    NOTEBOOK-012

Kullanım için README.md dosyasına bakın.
"""

import argparse
import os
import platform
import re
import socket
import subprocess
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.exit("HATA: 'openpyxl' paketi kurulu değil.\n"
             "      Kurmak için: pip install -r requirements.txt")


# ==========================================================================
#  ETİKET AYARLARI -- yerleşimi değiştirmek için sadece bu bölümü düzenleyin
#  Tüm ölçüler milimetre (mm). Yazıcı 203 dpi = 8 nokta/mm.
# ==========================================================================

DPMM = 8                     # 203 dpi => 8. 300 dpi yazıcıda 12 yapın.

ETIKET_GENISLIK_MM = 50.0
ETIKET_YUKSEKLIK_MM = 30.0

KENAR_BOSLUK_X_MM = 2.0      # Sol boşluk (ve sağda bırakılacak pay)

# --- Etiket satırları -------------------------------------------------
# Her satır: (şablon, üst_konum_mm, font_yüksekliği_mm)
# Şablonda kullanılabilecek alanlar:
#   {sartname}  sheet adının başındaki numara (örn. 3, 4.1)
#   {seri}      Seri No sütunu
#   {urun}      ürün etiketi: Ürün id sütunu varsa o, yoksa KOD-{sıra}
#   {sira}      sıra no (SIRA_BASAMAK kadar sıfırla doldurulmuş)
#   {sira_ham}  sıra no, dolgusuz
#   {kod}       ÜRÜN_KODLARI tablosundaki kısa kod
#   {isim}      sheet adı (baştaki numara atılmış hâli)
ETIKET_SATIRLARI = [
    ("ŞARTNAME NO:{sartname}",  3.5, 6.4),
    ("S/N:{seri}",             11.75, 6.4),
    ("{urun}",                 20.0, 6.4),
]

# True: üç satır da aynı puntoda basılır (uzun bir seri no hepsini küçültür).
# False: her satır kendi genişliğine göre en büyük puntoda basılır.
SATIRLARI_ESITLE = False

# True: metin + barkod bloğu etikette dikey olarak ortalanır; yukarıdaki
# "üst konum" değerleri yalnızca satırların birbirine göre aralığını belirler.
# False: üst konumlar etiketin üstünden mutlak ölçülür.
DIKEY_ORTALA = True

SIRA_BASAMAK = 3             # 12 -> "012". 0 yaparsanız dolgu yapılmaz.

# --- Barkod (örnek etikette yok; istenirse açılır) --------------------
# Barkod/QR açıkken metin satırları ETIKET_SATIRLARI yerine aşağıdaki
# sıkışık yerleşimden okunur, böylece barkodun üstüne binmez.
BARKOD_GOSTER = False        # True yapın ya da --barkod ile çalıştırın

ETIKET_SATIRLARI_BARKODLU = [
    ("ŞARTNAME NO:{sartname}",  1.2, 4.6),
    ("S/N:{seri}",              6.5, 4.6),
    ("{urun}",                 11.8, 4.6),
]
BARKOD_Y_MM = 18.0
BARKOD_YUKSEKLIK_MM = 9.5
BARKOD_MODUL = 0             # Çizgi kalınlığı (nokta): 0 = otomatik (sığan en geniş), 1-4 sabit
BARKOD_MODUL_TERCIHI = (3, 2, 1)   # Otomatik modda denenecek sıra
BARKOD_HRI = False

QR_BUYUKLUK = 4              # --qr için ^BQ büyütme oranı (1-10)
QR_X_MM = 37.0
QR_Y_MM = 19.0

# --- Genel ------------------------------------------------------------
MIN_FONT_MM = 2.0            # Otomatik küçültmede alt sınır
YAZDIRMA_HIZI = 3            # ^PR (ips), 2..6. Yavaş = daha net
KARARTMA = None              # ^MD, -30..30. None = yazıcının kendi ayarı.
                             # Barkod çizgileri kalın/yayılmış çıkıyorsa düşürün (örn. -5)

# --- Dijital önizleme (--onizleme) ------------------------------------
# Yazıcı olmadan etiketin nasıl çıkacağını PNG olarak gösterir. Tamamen
# çevrimdışıdır, hiçbir veri dışarı gönderilmez. Pillow paketi gerekir.
ONIZLEME_OLCEK = 4           # 1 yazıcı noktası kaç piksel çizilsin
ONIZLEME_MAKS = 12           # Tek PNG'de en fazla kaç etiket gösterilsin
ONIZLEME_SUTUN = 3           # Önizleme ızgarasında sütun sayısı
# Yazıcının yerleşik fontu CG Triumvirate Bold Condensed'tır. Önizlemede
# ona en yakın sistem fontu kullanılır; ilk bulunan seçilir.
ONIZLEME_FONTLARI = [
    "/System/Library/Fonts/Supplemental/Arial Narrow Bold.ttf",
    "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
    "C:/Windows/Fonts/arialnb.ttf",
    "C:/Windows/Fonts/arialbd.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed-Bold.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
]

# Sheet başına ürün kısa kodu. Ürün id sütunu dolu olan sheet'lerde
# bu tablo kullanılmaz, Excel'deki Ürün id değeri aynen basılır.
# Anahtarlar sheet adının normalize edilmiş hâlidir (küçük harf, Türkçe sade).
URUN_KODLARI = {
    "3.hp elitebook 8g2i 14 ai": "NOTEBOOK",
    "1. workstation": "WORKSTATION",
    "11.monitor": "MONITOR",
    "4. akilli telefon": "PHONE",
    "4.1 sarj adaptoru": "CHARGE ADAPTER",
    # Aşağıdakiler için kod bilinmiyor; doldurulmazsa sheet adı kullanılır
    # ve script uyarı verir:
    # "15. yazici lexmark mx532":      "PRINTER",
    # "17. a0 plotter epson surecolor": "PLOTTER",
    # "18.zebra barkod okuyucu":       "BARCODE SCANNER",
    # "19. web kamerasi logitech":     "WEBCAM",
    # "20. tripod kameralar icin":     "TRIPOD",
    # "23. usbcto6porthub":            "USB HUB",
    # "39.mouse kablosuz":             "MOUSE",
    # "6. logitech rallybar mini":     "RALLYBAR",
    # "21. hp usb-c to rj45 adapter g2": "RJ45 ADAPTER",
}

# ==========================================================================
#  Buradan sonrasını değiştirmeniz gerekmez
# ==========================================================================

def olculeri_hesapla():
    """mm cinsinden sabitlerden nokta (dot) değerlerini türet.
    Modül yüklenirken bir kez çağrılır; GUI boyutları değiştirdiğinde tekrar çağırır."""
    global GENISLIK_DOT, YUKSEKLIK_DOT, SOL_DOT, KULLANILABILIR_DOT
    GENISLIK_DOT = int(round(ETIKET_GENISLIK_MM * DPMM))
    YUKSEKLIK_DOT = int(round(ETIKET_YUKSEKLIK_MM * DPMM))
    SOL_DOT = int(round(KENAR_BOSLUK_X_MM * DPMM))
    KULLANILABILIR_DOT = GENISLIK_DOT - 2 * SOL_DOT


olculeri_hesapla()

# Yazıcının yerleşik ^A0 fontunda (CG Triumvirate Bold Condensed) her
# karakterin ilerlemesi, punto yüksekliğine oranla. Arial Narrow Bold'dan
# ölçüldü ve gerçek baskıyla doğrulandı (%2-7 sapma, güvenli tarafta).
# Punto küçültme kararı bu tabloyla verilir; taşmayı ayrıca ^FB keser.
KARAKTER_GENISLIK = {
    " ": 0.228, "#": 0.456, "(": 0.273, ")": 0.273, "*": 0.319, "+": 0.479,
    ",": 0.228, "-": 0.273, ".": 0.228, "/": 0.228, "0": 0.456, "1": 0.456,
    "2": 0.456, "3": 0.456, "4": 0.456, "5": 0.456, "6": 0.456, "7": 0.456,
    "8": 0.456, "9": 0.456, ":": 0.273, "A": 0.592, "B": 0.592, "C": 0.592,
    "D": 0.592, "E": 0.547, "F": 0.501, "G": 0.638, "H": 0.592, "I": 0.228,
    "J": 0.456, "K": 0.592, "L": 0.501, "M": 0.683, "N": 0.592, "O": 0.638,
    "P": 0.547, "Q": 0.638, "R": 0.592, "S": 0.547, "T": 0.501, "U": 0.592,
    "V": 0.547, "W": 0.774, "X": 0.547, "Y": 0.547, "Z": 0.501, "_": 0.456,
    "a": 0.456, "b": 0.501, "c": 0.456, "d": 0.501, "e": 0.456, "f": 0.273,
    "g": 0.501, "h": 0.501, "i": 0.228, "j": 0.228, "k": 0.456, "l": 0.228,
    "m": 0.729, "n": 0.501, "o": 0.501, "p": 0.501, "q": 0.501, "r": 0.319,
    "s": 0.456, "t": 0.273, "u": 0.501, "v": 0.456, "w": 0.638, "x": 0.456,
    "y": 0.456, "z": 0.410, "Ç": 0.592, "Ö": 0.638, "Ü": 0.592, "ç": 0.456,
    "ö": 0.501, "ü": 0.501, "Ğ": 0.638, "ğ": 0.501, "İ": 0.228, "ı": 0.228,
    "Ş": 0.547, "ş": 0.456,
}
KARAKTER_GENISLIK_VARSAYILAN = 0.60   # Tabloda olmayan karakterler için
GUVENLIK_PAYI = 1.03                   # Hesaplanan genişliğe eklenen pay

TR_ASCII = str.maketrans({
    "ş": "s", "Ş": "S", "ğ": "g", "Ğ": "G", "ı": "i", "İ": "I",
    "ö": "o", "Ö": "O", "ü": "u", "Ü": "U", "ç": "c", "Ç": "C",
    "â": "a", "Â": "A", "î": "i", "Î": "I", "û": "u", "Û": "U",
})

BASLIK_PARCALARI = ("seri no", "s/n", "s.no", "s/no", "sira no", "pc no",
                    "urun id", "sartname", "palet")
SERI_BASLIKLARI = ("seri no", "s/n", "s/no", "seri numarasi")
# Dikkat: "S.NO" (nokta) sıra no demektir, seri no değil -- bkz. sutunlari_bul()


def mm(v):
    return int(round(v * DPMM))


def normalize(s):
    """Karşılaştırma için: Türkçe sadeleştir, küçük harfe indir, boşluk sıkıştır."""
    if s is None:
        return ""
    s = str(s).strip().translate(TR_ASCII).lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s)


def hucre_metin(v):
    """Hücreyi metne çevir: 1234.0 ve bilimsel gösterim olmadan, sıfırlar korunarak."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def zpl_guvenli(s):
    return s.replace("^", "-").replace("~", "-")


# --------------------------------------------------------------------------
#  Excel okuma
# --------------------------------------------------------------------------

class Etiket:
    __slots__ = ("sira", "seri", "urun", "kod", "isim", "sartname",
                 "urun_id", "satir", "sheet", "index")


def baslik_satiri_mi(ws, r):
    hits = 0
    for c in range(1, ws.max_column + 1):
        n = normalize(ws.cell(r, c).value)
        if not n:
            continue
        if n == "id" or any(p in n for p in BASLIK_PARCALARI):
            hits += 1
    return hits >= 2


def sutunlari_bul(ws, hdr_row):
    seri = sira = urun = None
    adaylar = []
    for c in range(1, ws.max_column + 1):
        n = normalize(ws.cell(hdr_row, c).value)
        if not n:
            continue
        if any(p in n for p in SERI_BASLIKLARI):
            adaylar.append(c)
            if seri is None:
                seri = c
        elif sira is None and n in ("sira no", "pc no", "s.no", "s no",
                                    "sira", "no", "sira numarasi"):
            sira = c
        elif urun is None and (n == "id" or "urun id" in n):
            urun = c
    return seri, sira, urun, adaylar


def sartname_no(title):
    m = re.match(r"\s*(\d+(?:\.\d+)*)", str(title))
    return m.group(1) if m else ""


def sheet_adi_temizle(title):
    """'4.1 ŞARJ ADAPTÖRÜ' -> 'ŞARJ ADAPTÖRÜ'"""
    t = re.sub(r"^\s*\d+(\.\d+)*\s*\.?\s*", "", str(title)).strip()
    return t or str(title).strip()


def urun_etiketi(kod, sira):
    """'NOTEBOOK' + '12' -> 'NOTEBOOK-012'"""
    return "%s-%s" % (kod, sira.zfill(SIRA_BASAMAK) if SIRA_BASAMAK else sira)


def urun_kodu(title, uyarilar, zorla=None):
    """Sheet için ürün kısa kodu. 'zorla' verilirse (--kod) tablo okunmaz."""
    if zorla:
        return zorla
    n = normalize(title)
    if n in URUN_KODLARI:
        return URUN_KODLARI[n]
    kod = sheet_adi_temizle(title).upper()
    uyarilar.append("ürün kısa kodu tanımlı değil, sheet adı kullanıldı: %r\n"
                    "    (--kod KISAKOD ile komut satırından verebilir ya da\n"
                    "     etiket_bas.py içindeki URUN_KODLARI tablosuna "
                    "\"%s\": \"KOD\" satırını ekleyebilirsiniz)" % (kod, n))
    return kod


def sheet_oku(ws, kod_zorla=None):
    """Bir sheet'i oku. Dönüş: (etiketler, uyarılar).

    kod_zorla verilirse (--kod) ürün satırı, Excel'deki 'Ürün id' sütunu dolu
    olsa bile KOD-sıra biçiminde üretilir.
    """
    uyarilar = []
    hdr_row = None
    for r in range(1, min(ws.max_row, 10) + 1):
        if baslik_satiri_mi(ws, r):
            hdr_row = r
            break
    if hdr_row is None:
        return [], ["başlık satırı bulunamadı"]

    seri_c, sira_c, urun_c, adaylar = sutunlari_bul(ws, hdr_row)
    if seri_c is None:
        return [], ["'Seri No' sütunu bulunamadı"]
    if len(adaylar) > 1:
        uyarilar.append("birden fazla seri no sütunu var (sütun %s), ilki kullanıldı (%d)"
                        % (", ".join(map(str, adaylar)), seri_c))
    if sira_c is None:
        uyarilar.append("'Sıra No' başlığı yok, sıra numarası sayılarak üretildi")

    sart = sartname_no(ws.title)
    if not sart:
        uyarilar.append("sheet adında şartname numarası yok, boş bırakıldı")
    isim = sheet_adi_temizle(ws.title)
    kod = urun_kodu(ws.title, uyarilar, zorla=kod_zorla)

    etiketler = []
    idx = 0
    onceki = None
    for r in range(hdr_row + 1, ws.max_row + 1):
        if baslik_satiri_mi(ws, r):
            seri_c, sira_c, urun_c, adaylar = sutunlari_bul(ws, r)
            uyarilar.append("satır %d: ikinci başlık satırı -> yeni tablo bloğu" % r)
            if seri_c is None:
                uyarilar.append("satır %d sonrasında seri no sütunu yok, durduruldu" % r)
                break
            onceki = None
            continue

        seri = hucre_metin(ws.cell(r, seri_c).value)
        if not seri:
            continue
        ns = normalize(seri)
        if ns == "id" or any(p in ns for p in SERI_BASLIKLARI):
            uyarilar.append("satır %d: başlık metni gibi görünen değer atlandı (%r)" % (r, seri))
            continue

        sira_raw = hucre_metin(ws.cell(r, sira_c).value) if sira_c else ""
        if sira_raw:
            try:
                n = int(float(sira_raw))
                if onceki is not None and n < onceki:
                    uyarilar.append("satır %d: sıra no %d -> %d (geri gitti, "
                                    "yeni blok olabilir)" % (r, onceki, n))
                onceki = n
                sira_raw = str(n)
            except ValueError:
                pass
            sira = sira_raw
        else:
            sira = str(idx + 1)

        urun_id = hucre_metin(ws.cell(r, urun_c).value) if urun_c else ""
        idx += 1

        e = Etiket()
        e.sira = sira
        e.seri = seri
        e.kod = kod
        e.isim = isim
        e.sartname = sart
        e.urun_id = urun_id
        e.satir = r
        e.sheet = ws.title
        # --kod verildiyse Excel'deki Ürün id yok sayılır, kod-sıra üretilir
        e.urun = (urun_etiketi(kod, sira) if (kod_zorla or not urun_id)
                  else urun_id)
        etiketler.append(e)
        e.index = idx

    return etiketler, uyarilar


def excel_yukle(yol):
    if not os.path.exists(yol):
        sys.exit("HATA: Excel dosyası bulunamadı: %s" % yol)
    if not yol.lower().endswith((".xlsx", ".xlsm")):
        sys.exit("HATA: Desteklenmeyen dosya türü (.xlsx veya .xlsm olmalı): %s" % yol)
    try:
        return openpyxl.load_workbook(yol, data_only=True)
    except Exception as e:
        sys.exit("HATA: Excel dosyası açılamadı: %s\n      %s" % (yol, e))


# --------------------------------------------------------------------------
#  ZPL üretimi
# --------------------------------------------------------------------------

def metin_genisligi(metin, h):
    """Metnin ^A0N,h,h ile basıldığında kaplayacağı genişlik (dot)."""
    birim = sum(KARAKTER_GENISLIK.get(c, KARAKTER_GENISLIK_VARSAYILAN) for c in metin)
    return birim * h * GUVENLIK_PAYI


def font_sigdir(metin, genislik_dot, istenen_mm):
    """Metni verilen genişliğe sığdıran en büyük font yüksekliğini (dot) döndür."""
    h = mm(istenen_mm)
    if not metin:
        return h
    birim = metin_genisligi(metin, 1)
    if birim > 0:
        h = min(h, int(genislik_dot / birim))
    return max(h, mm(MIN_FONT_MM))


def code128_modul_sayisi(n):
    """Code 128 subset B, en kötü durum modül sayısı."""
    return 11 * (n + 2) + 13


def barkod_payload(seri, uyarilar):
    """Code 128 yalnızca ASCII 32-126 kodlar."""
    temiz = "".join(c for c in seri.translate(TR_ASCII) if 32 <= ord(c) <= 126)
    if temiz != seri:
        uyarilar.append("seri no %r barkodda %r olarak kodlandı (ASCII dışı karakter)"
                        % (seri, temiz))
    return temiz


def etiket_satirlari(e, tr_ascii=False, barkodlu=False):
    """Etiket satırı şablonlarını doldur, punto hesapla.
    ZPL üretimi ve önizleme bu aynı fonksiyonu kullanır.
    Dönüş: [(metin, üst_konum_mm, punto_dot), ...]"""
    duzen = ETIKET_SATIRLARI_BARKODLU if barkodlu else ETIKET_SATIRLARI
    alanlar = {
        "sartname": e.sartname,
        "seri": e.seri,
        "urun": e.urun,
        "sira": e.sira.zfill(SIRA_BASAMAK) if SIRA_BASAMAK else e.sira,
        "sira_ham": e.sira,
        "kod": e.kod,
        "isim": e.isim,
        "urun_id": e.urun_id,
    }
    satirlar = []
    for sablon, y_mm, font_mm in duzen:
        try:
            metin = sablon.format(**alanlar)
        except KeyError as k:
            sys.exit("HATA: etiket satırı şablonunda tanımsız alan: %s\n"
                     "      Kullanılabilir alanlar: %s"
                     % (k, ", ".join("{%s}" % a for a in alanlar)))
        if not metin:
            continue
        metin = zpl_guvenli(metin)
        if tr_ascii:
            metin = metin.translate(TR_ASCII)
        satirlar.append((metin, y_mm, font_sigdir(metin, KULLANILABILIR_DOT, font_mm)))

    if SATIRLARI_ESITLE and satirlar:
        ortak = min(h for _, _, h in satirlar)
        satirlar = [(m, y, ortak) for m, y, _ in satirlar]
    return satirlar


def barkod_modulu(payload, uyarilar):
    """Code 128 çizgi kalınlığı (dot). Sığmazsa None."""
    moduller = code128_modul_sayisi(len(payload))
    if BARKOD_MODUL:
        if moduller * BARKOD_MODUL <= KULLANILABILIR_DOT:
            return BARKOD_MODUL
        uyarilar.append("seri no %r için %d noktalık barkod etikete sığmıyor, "
                        "otomatik kalınlık kullanıldı" % (payload, BARKOD_MODUL))
    return next((x for x in BARKOD_MODUL_TERCIHI if moduller * x <= KULLANILABILIR_DOT), None)


def etiket_yerlesim(e, tr_ascii=False, barkod=False, qr=False, uyarilar=None):
    """Bir etiketin tüm öğelerini nokta (dot) koordinatlarıyla hesapla.
    DIKEY_ORTALA açıksa blok etikete ortalanır. ZPL ve önizleme bunu kullanır.

    Dönüş: dict(satirlar=[(metin, y_dot, h_dot)], payload, barkod_w, barkod_y,
                 barkod_h, qr_y, qr_x)
    """
    uyarilar = uyarilar if uyarilar is not None else []
    satirlar = [(m, mm(y), h) for m, y, h in etiket_satirlari(e, tr_ascii, barkodlu=(barkod or qr))]
    payload = barkod_payload(e.seri, uyarilar) if (barkod or qr) else ""
    barkod_w = barkod_modulu(payload, uyarilar) if (barkod and payload) else None
    if barkod and payload and barkod_w is None:
        uyarilar.append("seri no %r (%d karakter) barkod olarak sığmıyor, barkod basılmadı"
                        % (e.seri, len(payload)))
    barkod_y, barkod_h = mm(BARKOD_Y_MM), mm(BARKOD_YUKSEKLIK_MM)
    qr_y, qr_x = mm(QR_Y_MM), mm(QR_X_MM)

    if DIKEY_ORTALA:
        ustler = [y for _, y, _ in satirlar]
        altlar = [y + h for _, y, h in satirlar]
        if barkod_w:
            ustler.append(barkod_y); altlar.append(barkod_y + barkod_h)
        if qr and payload:
            ustler.append(qr_y); altlar.append(qr_y + 33 * QR_BUYUKLUK)   # ~25 modül + sessiz bölge
        if ustler:
            ust, alt = min(ustler), max(altlar)
            kaydir = (YUKSEKLIK_DOT - (alt - ust)) // 2 - ust
            kaydir = max(kaydir, -ust)                      # üstten taşmasın
            satirlar = [(m, y + kaydir, h) for m, y, h in satirlar]
            barkod_y += kaydir
            qr_y += kaydir

    return dict(satirlar=satirlar, payload=payload, barkod_w=barkod_w,
                barkod_y=barkod_y, barkod_h=barkod_h, qr_y=qr_y, qr_x=qr_x)


def etiket_zpl(e, kopya=1, barkod=False, qr=False, tr_ascii=False, uyarilar=None):
    uyarilar = uyarilar if uyarilar is not None else []
    Y = etiket_yerlesim(e, tr_ascii, barkod, qr, uyarilar)
    out = ["^XA", "^CI28",
           "^PW%d" % GENISLIK_DOT, "^LL%d" % YUKSEKLIK_DOT, "^LH0,0", "^LS0"]

    for metin, y, h in Y["satirlar"]:
        out.append("^FO%d,%d^A0N,%d,%d^FB%d,1,0,L,0^FD%s^FS"
                   % (SOL_DOT, y, h, h, KULLANILABILIR_DOT, metin))

    if qr and Y["payload"]:
        out.append("^FO%d,%d^BQN,2,%d,M,7^FDMM,A%s^FS"
                   % (Y["qr_x"], Y["qr_y"], QR_BUYUKLUK, zpl_guvenli(Y["payload"])))
    elif Y["barkod_w"]:
        out.append("^BY%d,3,%d" % (Y["barkod_w"], Y["barkod_h"]))
        out.append("^FO%d,%d^BCN,%d,%s,N,N,A^FD%s^FS"
                   % (SOL_DOT, Y["barkod_y"], Y["barkod_h"],
                      "Y" if BARKOD_HRI else "N", zpl_guvenli(Y["payload"])))

    if kopya > 1:
        out.append("^PQ%d,0,0,N" % kopya)
    out.append("^XZ")
    return "".join(out) + "\n"


def is_basligi():
    """İşin başında bir kez gönderilen yazıcı ayarları."""
    p = ["^XA", "^CI28", "^PW%d" % GENISLIK_DOT, "^LL%d" % YUKSEKLIK_DOT,
         "^LH0,0", "^LS0", "^MNY", "^PR%d" % YAZDIRMA_HIZI]
    if KARARTMA is not None:
        p.append("^MD%d" % KARARTMA)
    p.append("^XZ")
    return "".join(p) + "\n"


def is_uret(etiketler, kopya=1, barkod=False, qr=False, tr_ascii=False):
    uyarilar = []
    parcalar = [is_basligi()]
    for e in etiketler:
        parcalar.append(etiket_zpl(e, kopya=kopya, barkod=barkod, qr=qr,
                                   tr_ascii=tr_ascii, uyarilar=uyarilar))
    return "".join(parcalar), uyarilar


# --------------------------------------------------------------------------
#  Dijital önizleme (yazıcı olmadan PNG)
# --------------------------------------------------------------------------

def _onizleme_font(boyut_px):
    from PIL import ImageFont
    for yol in ONIZLEME_FONTLARI:
        if os.path.exists(yol):
            try:
                return ImageFont.truetype(yol, boyut_px)
            except Exception:
                continue
    return ImageFont.load_default()


def _onizleme_font_adi():
    for yol in ONIZLEME_FONTLARI:
        if os.path.exists(yol):
            return os.path.basename(yol)
    return "PIL varsayılan fontu"


def _code128b_desen(veri):
    """Code 128-B bar/boşluk genişlikleri. Önizleme için; yazıcı kendi kodlar."""
    P = ("212222 222122 222221 121223 121322 131222 122213 122312 132212 221213 "
         "221312 231212 112232 122132 122231 113222 123122 123221 223211 221132 "
         "221231 213212 223112 312131 311222 321122 321221 312212 322112 322211 "
         "212123 212321 232121 111323 131123 131321 112313 132113 132311 211313 "
         "231113 231311 112133 112331 132131 113123 113321 133121 313121 211331 "
         "231131 213113 213311 213131 311123 311321 331121 312113 312311 332111 "
         "314111 221411 431111 111224 111422 121124 121421 141122 141221 112214 "
         "112412 122114 122411 142112 142211 241211 221114 413111 241112 134111 "
         "111242 121142 121241 114212 124112 124211 411212 421112 421211 212141 "
         "214121 412121 111143 111341 131141 114113 114311 411113 411311 113141 "
         "114131 311141 411131 211412 211214 211232 2331112").split()
    kodlar = [104]                                   # START B
    for ch in veri:
        kodlar.append(ord(ch) - 32)
    toplam = kodlar[0] + sum(k * (i + 1) for i, k in enumerate(kodlar[1:]))
    kodlar.append(toplam % 103)                      # kontrol
    kodlar.append(106)                               # STOP
    return "".join(P[k] for k in kodlar)


def etiket_ciz(d, e, cx, cy, k, barkod=False, qr=False, tr_ascii=False):
    """Tek bir etiketi (cx, cy) köşesinden başlayarak PIL çizim nesnesine çiz.
    k = ölçek (1 yazıcı noktası kaç piksel). ZPL ile aynı yerleşimi kullanır."""
    et_w, et_h = GENISLIK_DOT * k, YUKSEKLIK_DOT * k
    d.rectangle([cx, cy, cx + et_w, cy + et_h], fill=(255, 255, 255),
                outline=(150, 155, 162))
    # kenar boşluğu kılavuzu
    d.rectangle([cx + SOL_DOT * k, cy, cx + et_w - SOL_DOT * k, cy + et_h],
                outline=(232, 234, 238))

    Y = etiket_yerlesim(e, tr_ascii, barkod, qr)
    for metin, y, h in Y["satirlar"]:
        f = _onizleme_font(max(1, int(h * k)))
        d.text((cx + SOL_DOT * k, cy + y * k), metin, font=f, fill=(0, 0, 0))

    if qr and Y["payload"]:
        qw = int(25 * QR_BUYUKLUK * k)
        x0, y0 = cx + Y["qr_x"] * k, cy + Y["qr_y"] * k
        d.rectangle([x0, y0, x0 + qw, y0 + qw], outline=(0, 0, 0))
        d.text((x0 + 3, y0 + 3), "QR", font=_onizleme_font(max(10, 6 * k)), fill=(120, 120, 120))
    elif Y["barkod_w"]:
        x = cx + SOL_DOT * k
        y0 = cy + Y["barkod_y"] * k
        y1 = y0 + Y["barkod_h"] * k
        siyah = True
        for ch in _code128b_desen(Y["payload"]):
            gen = int(ch) * Y["barkod_w"] * k
            if siyah:
                d.rectangle([x, y0, x + gen - 1, y1], fill=(0, 0, 0))
            x += gen
            siyah = not siyah


def onizleme_goruntu(etiketler, barkod=False, qr=False, tr_ascii=False,
                     olcek=None, sutun=None, altyazi_goster=True):
    """Etiketleri ızgara halinde çizip PIL Image döndür. Dosyaya yazmaz.
    GUI ve --onizleme bunu kullanır."""
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        sys.exit("HATA: Önizleme için 'Pillow' paketi gerekir.\n"
                 "      Kurmak için: pip install Pillow")

    k = olcek or ONIZLEME_OLCEK
    gosterilecek = etiketler[:ONIZLEME_MAKS]
    sutun = max(1, min(sutun or ONIZLEME_SUTUN, len(gosterilecek)))
    satir = (len(gosterilecek) + sutun - 1) // sutun

    et_w, et_h = GENISLIK_DOT * k, YUKSEKLIK_DOT * k
    bosluk = 5 * k
    altyazi = 8 * k if altyazi_goster else 0
    tuval = Image.new("RGB",
                      (sutun * et_w + (sutun + 1) * bosluk,
                       satir * (et_h + altyazi) + (satir + 1) * bosluk),
                      (225, 227, 230))
    d = ImageDraw.Draw(tuval)
    kucuk = _onizleme_font(max(10, altyazi - 2 * k)) if altyazi_goster else None

    for i, e in enumerate(gosterilecek):
        cx = bosluk + (i % sutun) * (et_w + bosluk)
        cy = bosluk + (i // sutun) * (et_h + altyazi + bosluk)
        etiket_ciz(d, e, cx, cy, k, barkod=barkod, qr=qr, tr_ascii=tr_ascii)
        if altyazi_goster:
            d.text((cx, cy + et_h + k),
                   "Excel satır %d  ·  etiket #%d" % (e.satir, e.index),
                   font=kucuk, fill=(90, 95, 102))
    return tuval


def onizleme_uret(etiketler, cikti, barkod=False, qr=False, tr_ascii=False):
    """--onizleme: görüntüyü üretip dosyaya yaz, özet bas."""
    tuval = onizleme_goruntu(etiketler, barkod=barkod, qr=qr, tr_ascii=tr_ascii)
    k = ONIZLEME_OLCEK
    gosterilecek = etiketler[:ONIZLEME_MAKS]

    try:
        tuval.save(cikti)
    except OSError as ex:
        sys.exit("HATA: Önizleme kaydedilemedi: %s\n      %s" % (cikti, ex))

    print("\nÖnizleme yazıldı -> %s" % os.path.abspath(cikti))
    print("  %d etiket çizildi (toplam %d seçili; ONIZLEME_MAKS ile artırılır)"
          % (len(gosterilecek), len(etiketler)))
    print("  Ölçek: 1 nokta = %d piksel  |  etiket %d x %d nokta" % (k, GENISLIK_DOT, YUKSEKLIK_DOT))
    print("  Font : %s" % _onizleme_font_adi())
    print("         (yazıcıdaki CG Triumvirate Bold Condensed'a yaklaşıktır;")
    print("          harf genişlikleri birebir aynı olmayabilir)")


# --------------------------------------------------------------------------
#  Yazıcı
# --------------------------------------------------------------------------

def yazicilari_listele():
    if platform.system() == "Windows":
        try:
            import win32print
        except ImportError:
            return []
        flags = win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS
        return [p[2] for p in win32print.EnumPrinters(flags, None, 1)]
    try:
        r = subprocess.run(["lpstat", "-a"], capture_output=True, text=True, timeout=10)
        return [ln.split()[0] for ln in r.stdout.splitlines() if ln.strip()]
    except Exception:
        return []


def yazici_sec():
    liste = yazicilari_listele()
    if not liste:
        sys.exit("HATA: Sistemde kurulu yazıcı bulunamadı.\n"
                 "      Yazıcıyı kurun veya --ip <adres> ile ağ üzerinden gönderin.\n"
                 "      (Windows'ta yazıcı listesi için 'pywin32' paketi gerekir.)")
    print("\nKurulu yazıcılar:")
    for i, p in enumerate(liste, 1):
        print("  %2d) %s" % (i, p))
    while True:
        s = input("Yazıcı numarası (iptal için Enter): ").strip()
        if not s:
            sys.exit("İptal edildi.")
        if s.isdigit() and 1 <= int(s) <= len(liste):
            return liste[int(s) - 1]
        print("Geçersiz numara.")


def gonder_windows(zpl, yazici):
    try:
        import win32print
    except ImportError:
        sys.exit("HATA: Windows'ta RAW baskı için 'pywin32' gerekir.\n"
                 "      Kurmak için: pip install pywin32")
    try:
        h = win32print.OpenPrinter(yazici)
    except Exception as e:
        sys.exit("HATA: Yazıcı bulunamadı veya açılamadı: %s\n      %s\n"
                 "      Kurulu yazıcıları görmek için: --printers" % (yazici, e))
    try:
        win32print.StartDocPrinter(h, 1, ("COP-31 Etiket", None, "RAW"))
        win32print.StartPagePrinter(h)
        win32print.WritePrinter(h, zpl.encode("utf-8"))
        win32print.EndPagePrinter(h)
        win32print.EndDocPrinter(h)
    finally:
        win32print.ClosePrinter(h)


def gonder_cups(zpl, yazici):
    try:
        r = subprocess.run(["lp", "-d", yazici, "-o", "raw", "-t", "COP-31 Etiket"],
                           input=zpl.encode("utf-8"), capture_output=True, timeout=180)
    except FileNotFoundError:
        sys.exit("HATA: 'lp' komutu bulunamadı (CUPS kurulu değil).")
    if r.returncode != 0:
        sys.exit("HATA: Yazıcıya gönderilemedi: %s\n      %s\n"
                 "      Kurulu yazıcıları görmek için: --printers"
                 % (yazici, r.stderr.decode("utf-8", "replace").strip()))


def gonder_ip(zpl, host, port):
    try:
        with socket.create_connection((host, port), timeout=15) as s:
            s.sendall(zpl.encode("utf-8"))
    except socket.timeout:
        sys.exit("HATA: Yazıcıya bağlanılamadı (zaman aşımı): %s:%d" % (host, port))
    except OSError as e:
        sys.exit("HATA: Yazıcıya bağlanılamadı: %s:%d\n      %s" % (host, port, e))


def gonder(zpl, yazici=None, ip=None, port=9100):
    if ip:
        gonder_ip(zpl, ip, port)
        return "%s:%d" % (ip, port)
    if not yazici:
        yazici = yazici_sec()
    if platform.system() == "Windows":
        gonder_windows(zpl, yazici)
    else:
        gonder_cups(zpl, yazici)
    return yazici


# --------------------------------------------------------------------------
#  Sheet seçimi ve listeleme
# --------------------------------------------------------------------------

def sheet_bul(wb, secim):
    adlar = wb.sheetnames
    if secim.isdigit():
        i = int(secim)
        if not 1 <= i <= len(adlar):
            sys.exit("HATA: Sheet numarası 1-%d aralığında olmalı, verilen: %d"
                     % (len(adlar), i))
        return adlar[i - 1]
    if secim in adlar:
        return secim
    hedef = normalize(secim)
    tam = [a for a in adlar if normalize(a) == hedef]
    if len(tam) == 1:
        return tam[0]
    kismi = [a for a in adlar if hedef in normalize(a)]
    if len(kismi) == 1:
        return kismi[0]
    if len(kismi) > 1:
        sys.exit("HATA: '%s' birden fazla sheet ile eşleşti:\n       %s"
                 % (secim, "\n       ".join(kismi)))
    sys.exit("HATA: '%s' adlı sheet bulunamadı. Listelemek için: --list" % secim)


def sheet_listesi_yaz(wb, detayli=False):
    print("\n %-3s %-36s %7s  %s" % ("No", "Sheet adı", "Etiket", "Excel satır (sıra no)"))
    print(" " + "-" * 88)
    toplam = bos = 0
    for i, ad in enumerate(wb.sheetnames, 1):
        et, uy = sheet_oku(wb[ad])
        toplam += len(et)
        if not et:
            bos += 1
            continue
        print(" %-3d %-36s %7d  %d-%d  (sıra %s-%s)"
              % (i, ad[:36], len(et), et[0].satir, et[-1].satir, et[0].sira, et[-1].sira))
        if detayli:
            for u in uy:
                print("     %s! %s" % (" " * 36, u))
    print(" " + "-" * 88)
    print(" %d sheet (%d tanesi boş, listelenmedi). Toplam %d etiket.\n"
          % (len(wb.sheetnames), bos, toplam))


def aralik_uygula(etiketler, bas, son, by_label, sheet_adi):
    alan = "etiket sırası" if by_label else "Excel satırı"
    key = (lambda e: e.index) if by_label else (lambda e: e.satir)
    sec = [e for e in etiketler
           if (bas is None or key(e) >= bas) and (son is None or key(e) <= son)]
    if not sec:
        sys.exit("HATA: '%s' sheet'inde verdiğiniz aralıkta etiket yok.\n"
                 "      İstenen: %s %s-%s\n"
                 "      Mevcut : %s %d-%d"
                 % (sheet_adi, alan,
                    bas if bas is not None else "baş", son if son is not None else "son",
                    alan, key(etiketler[0]), key(etiketler[-1])))
    return sec


# --------------------------------------------------------------------------
#  main
# --------------------------------------------------------------------------

def main():
    p = argparse.ArgumentParser(
        prog="etiket_bas.py",
        description="Excel'den Zebra ZD220 için 50x30 mm demirbaş etiketi basar.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""Örnekler:
  %(prog)s "COP-31 DEMİRBAŞLAR.xlsx" --list
  %(prog)s "COP-31 DEMİRBAŞLAR.xlsx" --sheet 1 --from 2 --to 51 --dry-run
  %(prog)s "COP-31 DEMİRBAŞLAR.xlsx" --sheet 1 --test
  %(prog)s "COP-31 DEMİRBAŞLAR.xlsx" --sheet 1 --from 2 --to 200 --printer ZDesigner
""")
    p.add_argument("excel", help="Excel dosyasının yolu")
    p.add_argument("--list", action="store_true",
                   help="Sheet'leri, etiket sayılarını ve satır aralıklarını listele")
    p.add_argument("--detay", action="store_true", help="--list çıktısında uyarıları da göster")
    p.add_argument("--sheet", help="Basılacak sheet: --list'teki numara veya sheet adı")
    p.add_argument("--from", dest="bas", type=int, help="Başlangıç Excel satır no (dahil)")
    p.add_argument("--to", dest="son", type=int, help="Bitiş Excel satır no (dahil)")
    p.add_argument("--by-label", action="store_true",
                   help="--from/--to Excel satırı değil, sheet içi etiket sırası (1..N)")
    p.add_argument("--kod", metavar="KISAKOD",
                   help="Ürün satırındaki kısa kodu elle belirle, örn. --kod HPLAPTOP "
                        "-> HPLAPTOP-001. Excel'deki 'Ürün id' sütunu varsa bile "
                        "bu kod kullanılır.")
    p.add_argument("--test", action="store_true", help="Sadece ilk etiketi bas")
    p.add_argument("--copies", type=int, default=1, help="Etiket başına kopya (varsayılan 1)")
    p.add_argument("--barkod", action="store_true", help="Seri noyu Code 128 barkod olarak da bas")
    p.add_argument("--qr", action="store_true", help="Seri noyu QR kod olarak da bas")
    p.add_argument("--tr-ascii", action="store_true",
                   help="Türkçe karakterleri ASCII'ye çevir (yazıcı fontu basamıyorsa)")
    p.add_argument("--printer", help="Yazıcı adı. Verilmezse kurulu yazıcılardan seçtirir.")
    p.add_argument("--printers", action="store_true", help="Kurulu yazıcıları listele")
    p.add_argument("--ip", help="Yazıcı IP adresi (RAW). --printer yerine kullanılır.")
    p.add_argument("--port", type=int, default=9100, help="RAW port (varsayılan 9100)")
    p.add_argument("--dry-run", action="store_true",
                   help="Yazıcıya gönderme; ZPL'i dosyaya yaz ve özeti göster")
    p.add_argument("--out", default="etiketler.zpl", help="--dry-run çıktı dosyası")
    p.add_argument("--onizleme", nargs="?", const="onizleme.png", metavar="PNG",
                   help="Yazıcıya göndermeden etiketin PNG önizlemesini üret "
                        "(varsayılan dosya: onizleme.png)")
    p.add_argument("--yes", "-y", action="store_true", help="Onay sorma")
    a = p.parse_args()

    if a.printers:
        liste = yazicilari_listele()
        print("Kurulu yazıcılar:" if liste else "Kurulu yazıcı bulunamadı.")
        for x in liste:
            print("  - %s" % x)
        return 0

    if a.copies < 1:
        sys.exit("HATA: --copies en az 1 olmalı.")
    if a.barkod and a.qr:
        sys.exit("HATA: --barkod ve --qr aynı anda kullanılamaz.")

    wb = excel_yukle(a.excel)

    if a.list:
        sheet_listesi_yaz(wb, detayli=a.detay)
        return 0

    interaktif = a.sheet is None
    if a.sheet:
        sheet_adi = sheet_bul(wb, a.sheet)
    else:
        sheet_listesi_yaz(wb)
        s = input("Basılacak sheet numarası veya adı (iptal için Enter): ").strip()
        if not s:
            print("İptal edildi.")
            return 1
        sheet_adi = sheet_bul(wb, s)

    etiketler, uyarilar = sheet_oku(wb[sheet_adi], kod_zorla=a.kod)
    if not etiketler:
        sys.exit("HATA: '%s' sheet'inde basılacak etiket yok.\n      %s"
                 % (sheet_adi, uyarilar[0] if uyarilar else "Seri no sütunu boş."))

    # Ürün satırının gerçek kaynağı (sheet_oku içindeki öncelik sırasıyla)
    id_li = sum(1 for e in etiketler if e.urun_id)
    if a.kod:
        kaynak = "--kod ile verildi"
    elif id_li == len(etiketler):
        kaynak = "Excel'deki 'Ürün id' sütunundan"
    elif id_li:
        kaynak = ("%d satırda Excel'deki 'Ürün id', kalan %d satırda %s-sıra"
                  % (id_li, len(etiketler) - id_li, etiketler[0].kod))
    elif normalize(sheet_adi) in URUN_KODLARI:
        kaynak = "URUN_KODLARI tablosundan"
    else:
        kaynak = "sheet adından türetildi"

    print("\nSheet            : %s" % sheet_adi)
    print("Şartname no      : %s" % (etiketler[0].sartname or "(yok)"))
    print("Ürün satırı      : %s  (%s)" % (etiketler[0].urun, kaynak))
    print("Geçerli etiket   : %d adet | Excel satır %d-%d | etiket sırası 1-%d"
          % (len(etiketler), etiketler[0].satir, etiketler[-1].satir, len(etiketler)))
    for u in uyarilar:
        print("  ! %s" % u)

    bas, son = a.bas, a.son
    if interaktif and bas is None and son is None:
        alan = "etiket sırası" if a.by_label else "Excel satır no"
        s1 = input("\nBaşlangıç %s (tümü için Enter): " % alan).strip()
        s2 = input("Bitiş %s (sona kadar için Enter): " % alan).strip()
        bas = int(s1) if s1.isdigit() else None
        son = int(s2) if s2.isdigit() else None

    if bas is not None and son is not None and bas > son:
        sys.exit("HATA: --from (%d), --to (%d) değerinden büyük olamaz." % (bas, son))

    secim = aralik_uygula(etiketler, bas, son, a.by_label, sheet_adi)
    if a.test:
        secim = secim[:1]
        print("\n--test: sadece ilk etiket basılacak.")

    zpl, zpl_uyari = is_uret(secim, kopya=a.copies, barkod=a.barkod, qr=a.qr,
                             tr_ascii=a.tr_ascii)

    alan = "etiket sırası" if a.by_label else "Excel satırı"
    key = (lambda e: e.index) if a.by_label else (lambda e: e.satir)
    print("\n" + "=" * 72)
    print("ÖZET")
    print("  Sheet            : %s" % sheet_adi)
    print("  Aralık           : %s %d - %d" % (alan, key(secim[0]), key(secim[-1])))
    print("  İlk etiket       : ŞARTNAME NO:%s | S/N:%s | %s"
          % (secim[0].sartname, secim[0].seri, secim[0].urun))
    print("  Son etiket       : ŞARTNAME NO:%s | S/N:%s | %s"
          % (secim[-1].sartname, secim[-1].seri, secim[-1].urun))
    print("  Etiket adedi     : %d  (kopya %d -> toplam %d)"
          % (len(secim), a.copies, len(secim) * a.copies))
    print("  Etiket boyutu    : %.1f x %.1f mm (%d x %d dot, %d dpmm)"
          % (ETIKET_GENISLIK_MM, ETIKET_YUKSEKLIK_MM, GENISLIK_DOT, YUKSEKLIK_DOT, DPMM))
    if a.barkod or a.qr:
        print("  Ek kod           : %s" % ("QR" if a.qr else "Code 128"))
    print("=" * 72)
    for u in dict.fromkeys(zpl_uyari):
        print("  ! %s" % u)

    if a.onizleme:
        onizleme_uret(secim, a.onizleme, barkod=a.barkod, qr=a.qr, tr_ascii=a.tr_ascii)
        if not a.dry_run:
            print("\nYazıcıya hiçbir şey gönderilmedi (--onizleme).")
            return 0

    if a.dry_run:
        try:
            with open(a.out, "w", encoding="utf-8", newline="\n") as f:
                f.write(zpl)
        except OSError as e:
            sys.exit("HATA: ZPL dosyası yazılamadı: %s\n      %s" % (a.out, e))
        print("\n--dry-run: yazıcıya hiçbir şey gönderilmedi.")
        print("ZPL yazıldı -> %s  (%d etiket, %d bayt)"
              % (os.path.abspath(a.out), len(secim), len(zpl.encode("utf-8"))))
        return 0

    if not a.yes:
        c = input("\n%d etiket basılacak. Devam edilsin mi? [e/H] " % (len(secim) * a.copies))
        if c.strip().lower() not in ("e", "evet", "y", "yes"):
            print("İptal edildi. Hiçbir etiket basılmadı.")
            return 1

    hedef = gonder(zpl, yazici=a.printer, ip=a.ip, port=a.port)
    print("\nTamam: %d etiket tek iş olarak gönderildi -> %s"
          % (len(secim) * a.copies, hedef))
    print("Yarıda kalırsa devam etmek için:  --sheet \"%s\" --from <satır> --to %d"
          % (sheet_adi, key(secim[-1])))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print("\nİptal edildi.")
        sys.exit(130)
